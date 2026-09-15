import os
import pickle
from flask import Flask, request, jsonify
import faiss
from sentence_transformers import SentenceTransformer
from openai import OpenAI
import tempfile
from config import *
# File readers
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import pytesseract
import cv2
import json
# Calculate PDF identity
import hashlib
import subprocess
from dotenv import load_dotenv
import re

# .env file loading
load_dotenv()

# ========================================================================
# PDF Conversion Function (LibreOffice used for .docx, .pptx, etc. to PDF) 
# ========================================================================
def convert_to_pdf(input_path):
    output_dir = os.path.dirname(input_path)
    libreoffice_path = "soffice"
    result = subprocess.call([
        libreoffice_path,
        '--headless',
        '--convert-to', 'pdf',
        input_path,
        '--outdir', output_dir
    ])

    base = os.path.splitext(input_path)[0]
    pdf_path = base + ".pdf"

    #  CRITICAL CHECK
    if result != 0 or not os.path.exists(pdf_path):
        raise Exception("PDF conversion failed")

    return pdf_path

# INIT 
app = Flask(__name__)

#API openai client
client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY"),
    base_url="https://openrouter.ai/api/v1"
)

# Tesseract OCR for image-based PDFs and images
pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

# EMBEDDING MODEL
embed_model = SentenceTransformer("all-MiniLM-L6-v2")
DIM = 384


# ========================================================================
# LOAD / CREATE FAISS
# ========================================================================

if os.path.exists(FAISS_PATH) and os.path.getsize(FAISS_PATH) > 0:
    try:
        index = faiss.read_index(FAISS_PATH)
        # Documents are stored separately in a pickle file (rb mode for binary read as Pickle is a binary format)
        with open(DOCS_PATH, "rb") as f:
            documents = pickle.load(f)
    except:
        index = faiss.IndexFlatL2(DIM)
        documents = []
else:
    index = faiss.IndexFlatL2(DIM)
    documents = []


# SAVE "wb" for binary write
def save_state():
    faiss.write_index(index, FAISS_PATH)
    with open(DOCS_PATH, "wb") as f:
        pickle.dump(documents, f)
        
# ========================================================================
# FILE → TEXT Extraction (PDF, DOCX, PPTX, TXT, Images)
# ========================================================================
def extract_text(path):
  
    ext = os.path.splitext(path)[1].lower()  # it extract the extension and convert it to lower case for comparison
    if ext == ".pdf":
        try:
             reader = PdfReader(path)  
        except Exception as e:
            print("PDF READ ERROR:", str(e))
            return ""
        
        text = ""

        for page in reader.pages:
            text += page.extract_text() or "" # Extract text from each page if available, otherwise add an empty string

        # CLEAN AFTER FULL EXTRACTION
        
        lines = text.split("\n")
        clean_lines = []

        for line in lines:
            line = line.strip()

            if not line:
                continue

            if line.isdigit():
                continue

            if not any(char.isalnum() for char in line): # Skip lines that do not contain any alphanumeric characters
                continue

            clean_lines.append(line)

        text = "\n".join(clean_lines)


        # IF TEXT EMPTY → USE OCR
        if text.strip():
            return text

        print("PDF has no text, using OCR...")

        try:
            import pdf2image
            # convert Pdf pages to Images
            images = pdf2image.convert_from_path(path)

            ocr_text = ""

            for img in images:
                text = pytesseract.image_to_string(img)
                ocr_text += text + "\n"

            print("OCR TEXT SAMPLE:", ocr_text[:200])

            return ocr_text

        except Exception as e:
            print("OCR ERROR:", str(e))
            return ""

    
    elif ext == ".docx":
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs])

    elif ext == ".pptx":
        prs = Presentation(path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text

    elif ext == ".txt":
        return open(path, "r", encoding="utf-8").read()

    elif ext in [".png", ".jpg", ".jpeg"]:
        try:
            img = cv2.imread(path) # Read the image using OpenCV that make us to preprocess the image for better OCR accuracy
            gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY) # Convert to grayscale for better OCR accuracy

            text = pytesseract.image_to_string(gray)

            return text

        except Exception as e:
            print("OCR ERROR:", str(e))
            return ""

    return ""

# ========================================================================
# KEYWORD EXTRACTION for Notes uploaed that helps in keyword searching 
# ========================================================================

def extract_keywords(text):
    """
    Extract important domain-specific keywords from document text.
    If AI keyword extraction fails, return an empty list so that
    the normal upload flow is not affected.
    """

    if not text or not text.strip():
        return []

    # Keep prompt reasonably small
    clean_text = text[:12000]

    prompt = f"""
You are a document keyword extraction system.

Extract the 10 to 15 MOST IMPORTANT keywords or short technical concepts
from the document.

RULES:
- Return ONLY valid JSON.
- JSON format must be:
{{
  "keywords": ["keyword1", "keyword2", "keyword3"]
}}
- Extract meaningful subject/domain concepts.
- Prefer technical terms, topics, theories, technologies,
  methods, algorithms, models, important terminology and named concepts.
- Use short words or short phrases.
- Do NOT return generic words such as:
  "notes", "chapter", "student", "university", "important",
  "introduction", "page", "document", "question".
- Do NOT return duplicate keywords.
- Keep keywords in their normal readable form.
- Maximum 15 keywords.

DOCUMENT:
{clean_text}
"""

    try:
        raw = call_openai(
            prompt,
            max_tokens=300,
            json_mode=True
        )

        print("KEYWORD AI RESPONSE:", raw)

        data = json.loads(raw)

        keywords = data.get("keywords", [])

        if not isinstance(keywords, list):
            return []

        # Clean and normalize
        cleaned = []

        for keyword in keywords:
            if not isinstance(keyword, str):
                continue

            keyword = keyword.strip()

            if not keyword:
                continue

            if keyword.lower() not in [
                k.lower() for k in cleaned
            ]:
                cleaned.append(keyword)

        return cleaned[:15]

    except Exception as e:
        print("KEYWORD EXTRACTION ERROR:", str(e))

        # IMPORTANT:
        # Keyword extraction must NEVER break normal upload.
        return []

# API create for Extract keyword (called from Node.js)
@app.route("/extract-keywords", methods=["POST"])
def extract_keywords_endpoint():
    file = request.files.get("file")
    if not file:
        return jsonify({
            "keywords": [],
            "error": "No file"
        }), 400

    # IMPORTANT:
    # Create a PRIVATE temporary file for keyword extraction.
    # Do NOT use UPLOAD_DIR because Node.js also uses /app/data/temp for the original uploaded file.
    
    file_ext = os.path.splitext(file.filename)[1] or ".pdf"
    temp_file = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=file_ext
    )
    temp_path = temp_file.name
    temp_file.close()
    try:
        file.save(temp_path)
        # Reuse existing PDF/OCR extraction
        text = extract_text(temp_path)
        if not text.strip():
            return jsonify({
                "keywords": [],
                "error": "No readable text found"
            }), 200
        keywords = extract_keywords(text)
        print("EXTRACTED KEYWORDS:", keywords)
        return jsonify({
            "keywords": keywords
        })
    except Exception as e:
        print("KEYWORD ENDPOINT ERROR:", str(e))
        # Do NOT make upload dependent on this
        return jsonify({
            "keywords": [],
            "error": str(e)
        }), 200

    finally:
        # Delete ONLY the private keyword-extraction temp file.
        # This will NOT delete Node's /app/data/temp uploaded file.
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass


# ========================================================================
# CHUNKING
# ========================================================================

def chunk_text(text):

    # re.sub() means finds a pattern and replaces it with a specified string. Here, it replaces multiple spaces or tabs with a single space.
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)

    # Split document into paragraphs 
    paragraphs = []
    for p in re.split(r'\n\s*\n', text): # Split on two or more newlines (with optional whitespace in between).
        if p.strip():
            paragraphs.append(p.strip())

    MAX_CHUNK_SIZE = 1200
    OVERLAP = 200 # The next chunk shares some context with the previous one (here 200 characters) to maintain continuity.

    chunks = []

    for paragraph in paragraphs:

        # Split paragraph into sentences
        sentences = re.split(r'(?<=[.!?])\s+', paragraph) # s+ (one or more whitespace)

        current_chunk = "" # This temporarily stores sentences until they reach the maximum size.

        for sentence in sentences:

            sentence = sentence.strip()

            if not sentence:
                continue

            # Add sentence if it fits
            if len(current_chunk) + len(sentence) + 1 <= MAX_CHUNK_SIZE:

                if current_chunk:
                    current_chunk += " " + sentence
                else:
                    current_chunk = sentence

            else:
                # Save completed chunk
                if current_chunk:
                    chunks.append(current_chunk.strip())

                # Start next chunk with overlap
                overlap_text = current_chunk[-OVERLAP:] # Take the last 200 characters of current_chunk.

                current_chunk = overlap_text + " " + sentence

        # Save remaining text (that not reach 1200 characters) as a chunk
        if current_chunk:
            chunks.append(current_chunk.strip())

    return chunks

def add_chunks(chunks, source, file_hash=None):

    # Prevent duplicate indexing
    if file_hash:
        for doc in documents:
            if doc.get("file_hash") == file_hash:
                print("FILE ALREADY INDEXED:", file_hash)
                return False

    clean_chunks = []

    for chunk in chunks:

        chunk = chunk.strip() # only for removing leading and trailing whitespace

        # Skip empty or very short chunks
        if not chunk:
            continue

        if len(chunk.split()) < 15: # Skip chunks with fewer than 15 words as very short pieces generally provide poor semantic context for retrieval.
            continue

        clean_chunks.append(chunk)

    if not clean_chunks: # If no valid chunks remain after cleaning, return False to indicate that nothing was added.
        return False

    # Create embeddings
    emb = embed_model.encode(
        clean_chunks,
        batch_size=8, # embedding model processes up to 8 chunks at a time in a batch.
        convert_to_numpy=True
    ).astype("float32")

    # Add vectors to FAISS
    index.add(emb)

    # Store chunk information as FAISS stores the vectors, but we also need the original text and metadata.
    for chunk in clean_chunks:
        documents.append({
            "text": chunk,
            "source": source,
            "file_hash": file_hash
        })

    # Persist FAISS + documents
    save_state()

    return True

# ========================================================================
# RAG SEARCH
# ========================================================================

def search(query, k=TOP_K, source=None, session_id=None):
    # docString is a multi-line string that describes the function's purpose, parameters, and behavior.
    """ 
    RAG retrieval.
    Cases:
    1. source provided:
       Search ONLY this exact indexed file (Path comes from PDF reader Source).

    2. session_id provided:
       Search ONLY files belonging to this chat session (File attached from Chat Session).

    3. Neither provided:
       Search globally across all indexed files (No file attached or not comes from Source).
    """
    if not documents:
        return []
    
    query = query.strip()
    if not query:
        return []

    scoped_docs = [] # temporary selection/filter
    if source:
        scoped_docs = [ # list comprehension to filter documents by source
            doc for doc in documents
            if doc.get("source") == source
        ]
        print("RAG SCOPE: SPECIFIC FILE")
        print("SOURCE:", source)

    # File Attached from Chat Session (like a temporary file uploaded for this session)
    elif session_id:
        session_prefix = session_id.rstrip("/") + "/"
        session_docs = [
            doc for doc in documents
            if doc.get("source", "").startswith(session_prefix)
        ]
        if session_docs:
            scoped_docs = session_docs
            print("RAG SCOPE: CHAT SESSION")
            print("SESSION:", session_id)
        else:
            scoped_docs = documents
            print("RAG SCOPE: GLOBAL")
            print("NO FILE IN SESSION → GLOBAL RAG")
    else:
        # Normal global RAG
        scoped_docs = documents
        print("RAG SCOPE: GLOBAL")

    if not scoped_docs:
        print("RAG: No documents found in selected scope.")
        return []

    #  Get FAISS vector IDs belonging to scope
    scoped_indices = []
    for i, doc in enumerate(documents): # i => index in FAISS, doc => document metadata
        if doc in scoped_docs:
            scoped_indices.append(i)
    if not scoped_indices:
        return []

    #  Embed question (to retrieve semantically similar chunks)
    q_emb = embed_model.encode(
        [query], # Encode this one question as a batch of one item.
        convert_to_numpy=True
    ).astype("float32")

    # Search FAISS
    search_k = min( # min helps to take not more than actual chunks in FAISS.
        max(k * 10, 20), # max gives which is maximum in btw them (if k*10 is less than 20 then 20 will be used)
        index.ntotal
    )

    # D = distances, I = indices of nearest neighbors    { for us "smaller L2 distance means greater similarity"}.
    D, I = index.search(q_emb, search_k)

    candidates = []
    for idx, dist in zip(I[0], D[0]): # zip() combines the IDs and distances.
        # Invalid FAISS index
        if idx < 0 or idx >= len(documents):
            continue

        if idx not in scoped_indices: #FAISS gives us candidate vector IDs,if that is not in our scoped_indices we skip it.
            continue

        candidates.append({
            "text": documents[idx]["text"],
            "source": documents[idx]["source"],
            "distance": float(dist)
        })

    # Sort by semantic similarity
    candidates.sort(
        key=lambda x: x["distance"] #Sort each candidate according to its distance value
    )

    # Remove duplicate chunk text
    final = []
    seen = set() # set contains only unique elements.
    for item in candidates:
        text = item["text"].strip()
        if not text:
            continue
        if text in seen:
            continue
        seen.add(text)
        final.append({
            "text": text,
            "source": item["source"]
        })
        if len(final) >= k:
            break
    return final

# ========================================================================
# OPENAI CALL
# ========================================================================

def call_openai(prompt, max_tokens=700, json_mode=False):
    params = {
        "model": "openai/gpt-4o-mini",
        "messages": [{"role": "user", "content": prompt}],
        "max_tokens": max_tokens,
        "temperature": 0.5 # How much freedom the model has when choosing the next token 
    }
    
    if json_mode:
        params["response_format"] = {"type": "json_object"}

    res = client.chat.completions.create(**params)

    return res.choices[0].message.content

# ========================================================================
# PROMPTS
# ========================================================================

def build_answer_prompt(context, question):
    return f"""
You are a helpful AI assistant.

Answer the question using ONLY the given context.

RULES:
- Give a COMPLETE answer (not short or not too Long)
- Explain clearly in simple language
- If multiple points exist → use bullet points
- Combine information from multiple chunks if needed
- Do NOT say "not found"
- Do NOT guess
- Do NOT use external knowledge if any relevant data find from context
- Answer should feel like a short explanation, not just a line

Context:
{context}

Question:
{question}

Answer:
"""

def build_quiz_prompt(context, previous_quiz=None):
    previous_quiz = previous_quiz or []

    previous_questions = "\n".join(
        f"- {q.get('question', '')}"
        for q in previous_quiz
        if isinstance(q, dict) and q.get("question")
    )

    regeneration_rules = ""

    if previous_questions:
        regeneration_rules = f"""
IMPORTANT REGENERATION RULES:
- This is a regeneration request.
- Generate completely NEW questions.
- DO NOT repeat any previous question.
- DO NOT simply rephrase a previous question.
- DO NOT ask the same question with different options.
- Try to test different concepts from the provided context.

PREVIOUS QUESTIONS:
{previous_questions}
"""

    return f"""
You are a strict JSON generator.

TASK:
Generate EXACTLY 5 multiple choice questions (MCQs).

CRITICAL RULES:
- Output ONLY valid JSON
- DO NOT write any text before or after JSON
- JSON must start with {{ and end with }}
- Always return complete JSON

RULES:
- Exactly 5 questions
- Each question must be unique
- Each question must have exactly 4 options
- answer must be an index: 0, 1, 2, or 3
- DO NOT include numbering inside options
- explanation must be clear
- Questions must be based ONLY on the provided context
- Cover different concepts from the context where possible

{regeneration_rules}

FORMAT:
{{
  "quiz": [
    {{
      "question": "Question text",
      "options": [
        "Option 1",
        "Option 2",
        "Option 3",
        "Option 4"
      ],
      "answer": 0,
      "topic": "General",
      "explanation": "Explanation of correct answer"
    }}
  ]
}}

Context:
{context[:1200]}
"""

# ========================================================================
# UPLOAD (RAG INDEXING)
# ========================================================================

@app.route("/upload", methods=["POST"])
def upload():
    path = None
    try:
        import hashlib
        file = request.files["file"]
        file_bytes = file.read()
        file_hash = hashlib.sha256(file_bytes).hexdigest()
        file.seek(0) # Reset file pointer to the beginning after reading for hash calculation

        path = os.path.join(UPLOAD_DIR, file.filename)
        file.save(path)

        text = extract_text(path)
        if not text.strip():
            return jsonify({"error": "No text extracted"})

        chunks = chunk_text(text)
        print("CHUNKS COUNT:", len(chunks)) 
        session_id = request.form.get("session_id", "global")
        source = f"{session_id}/{file.filename}"

        added = add_chunks(
            chunks,
            source=source,
            file_hash=file_hash
        )
        if not added:
            return jsonify({
                "status": "already indexed"
    })
        return jsonify({"status": "indexed", "chunks": len(chunks)})
    except Exception as e:
        print("UPLOAD ERROR:", str(e))
        return jsonify({"error": str(e)}), 500
    finally:
        if path and os.path.exists(path):
            os.remove(path)   # we delete the uploaded file because we need it only for store it in FAISS and pickle for RAG.
        
# ========================================================================
# Convert
# ========================================================================

@app.route("/convert", methods=["POST"])
def convert_file():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "No file"}), 400

    input_path = os.path.join(UPLOAD_DIR, file.filename)
    file.save(input_path)

    ext = os.path.splitext(input_path)[1].lower() # Get the file extension in lowercase for comparison

    try:
        if ext == ".pdf":
            result = os.path.basename(input_path)
        else:
            pdf_path = convert_to_pdf(input_path)
            result = os.path.basename(pdf_path)

        return jsonify({"pdfPath": result})

    except Exception as e:
        print("CONVERSION ERROR:", str(e))
        return jsonify({"error": str(e)}), 500

    finally:
        if os.path.exists(input_path):
            os.remove(input_path) # We delete the original uploaded file after conversion to PDF to save storage space.

# ========================================================================
# ASK (RAG)
# ========================================================================

@app.route("/ask-rag", methods=["POST"])
def ask_rag():
    try:
        data = request.get_json() # This reads the JSON sent by Node.js.

        question = data.get("question")
        source = data.get("source")
        session_id = data.get("session_id")
        if not question:
            return jsonify({
                "error": "Question required"
            }), 400
        # RAG SEARCH
        context_chunks = search(
            question,
            k=TOP_K,
            source=source,
            session_id=session_id
        )
        # No relevant RAG context
        if not context_chunks:
            return jsonify({
                "answer":
                "This information is not available in the indexed Notes House material.",
                "mode": "rag"
            })
        # Build context
        context = "\n\n".join(
            item["text"]
            for item in context_chunks
        )
        # Generate answer
        prompt = build_answer_prompt(
            context,
            question
        )
        answer = call_openai(
            prompt,
            max_tokens=700
        )
        return jsonify({
            "answer": answer,
            "mode": "rag"
        })

    except Exception as e:
        print(
            "ASK_RAG ERROR:",
            str(e)
        )
        return jsonify({
            "error": str(e)
        }), 500  
    
# ========================================================================
# ASK PDF (RAG)
# ========================================================================

@app.route("/ask-pdf", methods=["POST"])
def ask_pdf():
    file = request.files.get("file")
    question = request.form.get("question")
    if not file or not question:
        return jsonify({
            "error": "Missing file or question"
        }), 400
        
    temp_path = os.path.join(
        UPLOAD_DIR,
        file.filename
    )
    try:
        file.save(temp_path)
        # Calculate file hash
        with open(temp_path, "rb") as f:
            file_hash = hashlib.sha256(
                f.read()
            ).hexdigest()
        # Check whether this exact file is already indexed in FAISS
        existing_source = None
        for doc in documents:
            if doc.get("file_hash") == file_hash:
                existing_source = doc.get("source")
                break
        # ---------------------------------------------
        # CASE 1: Already indexed
        # IMPORTANT: Do NOT extract/chunk/embed/add again.
        # ---------------------------------------------
        if existing_source:
            print("PDF ALREADY INDEXED")
            print("USING EXISTING SOURCE:", existing_source)
            results = search(
                question,
                k=TOP_K,
                source=existing_source
            )
        # ---------------------------------------------
        # CASE 2: New file
        # Index it under this chat session.
        # ---------------------------------------------
        else:
            text = extract_text(temp_path)
            if not text.strip():
                return jsonify({
                    "error": "No text extracted"
                }), 400
            chunks = chunk_text(text)
            session_id = request.form.get(
                "session_id",
                "temp_session"
            )
            source = (
                f"{session_id.rstrip('/')}/"
                f"{file.filename}"
            )
            added = add_chunks(
                chunks,
                source=source,
                file_hash=file_hash
            )
            if not added:
                return jsonify({
                    "error": "File could not be indexed"
                }), 500
            print("NEW PDF INDEXED")
            print("SOURCE:", source)
            # Search ONLY this session's file
            results = search(
                question,
                k=TOP_K,
                session_id=session_id
            )
        # Build context
        relevant_chunks = [
            item["text"]
            for item in results
        ]
        if not relevant_chunks:
            return jsonify({
                "answer": "No relevant information was found in this file."
            })
        context = "\n\n".join(
            relevant_chunks
        )
        # Generate answer
        prompt = build_answer_prompt(
            context,
            question
        )
        answer = call_openai(
            prompt,
            max_tokens=700
        )
        return jsonify({
            "answer": answer,
            "mode": "rag"
        })
    except Exception as e:
        print("ASK PDF ERROR:", str(e))
        return jsonify({
            "error": str(e)
        }), 500
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass

# ========================================================================
# ASK FILE (RAG)
# ========================================================================
@app.route("/ask-file", methods=["POST"])
def ask_file():
    file = request.files.get("file")
    question = request.form.get("question")
    if not file or not question:
        return jsonify({
            "error": "Missing file or question"
        }), 400

    temp_path = os.path.join(
        UPLOAD_DIR,
        file.filename
    )
    try:
        file.save(temp_path)
        # Calculate file hash
        with open(temp_path, "rb") as f:
            file_hash = hashlib.sha256(
                f.read()
            ).hexdigest()
        # ---------------------------------------------
        # Check if already indexed
        # ---------------------------------------------
        existing_source = None
        for doc in documents:
            if doc.get("file_hash") == file_hash:
                existing_source = doc.get("source")
                break
        # ---------------------------------------------
        # Already indexed
        # ---------------------------------------------
        if existing_source:
            print("FILE ALREADY INDEXED")
            print("USING EXISTING SOURCE:", existing_source)

            results = search(
                question,
                k=TOP_K,
                source=existing_source
            )
        # ---------------------------------------------
        # New file
        # ---------------------------------------------
        else:
            text = extract_text(temp_path)
            if not text.strip():
                return jsonify({
                    "answer": "No readable content found"
                })
            chunks = chunk_text(text)
            session_id = request.form.get(
                "session_id",
                "temp_session"
            )
            source = (
                f"{session_id.rstrip('/')}/"
                f"{file.filename}"
            )
            added = add_chunks(
                chunks,
                source=source,
                file_hash=file_hash
            )
            if not added:
                return jsonify({
                    "answer": "File could not be indexed"
                }), 500
            print("NEW FILE INDEXED")
            print("SOURCE:", source)
            results = search(
                question,
                k=TOP_K,
                session_id=session_id
            )
        # Build context
        relevant_chunks = [
            item["text"]
            for item in results
        ]
        if not relevant_chunks:
            return jsonify({
                "answer":
                "No relevant information was found in this file."
            })
        context = "\n\n".join( # \n\n = press Enter twice The first \n moves to the next line, and the second creates the empty line.
            relevant_chunks
        )
        prompt = build_answer_prompt(
            context,
            question
        )
        answer = call_openai(
            prompt,
            max_tokens=700
        )
        return jsonify({
            "answer": answer,
            "mode": "rag"
        })
    except Exception as e:
        print("ASK FILE ERROR:", str(e))
        return jsonify({
            "error": str(e)
        }), 500
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass

# ========================================================================
# File Based Quiz Generation
# ========================================================================

@app.route("/generate-quiz-file", methods=["POST"])
def generate_quiz_file():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "No file"}), 400
    temp_path = os.path.join(UPLOAD_DIR, file.filename)
    try:
        file.save(temp_path)
        #  Extract text
        text = extract_text(temp_path)
        if not text.strip():
            return jsonify({
                "quiz": [],
                "error": "No text found"
            }), 400
        #  Get regeneration information
        regenerate = (
            request.form.get("regenerate", "false").lower()
            == "true"
        )
        previous_quiz_raw = request.form.get(
            "previousQuiz",
            "[]"
        )
        try:
            previous_quiz = json.loads(previous_quiz_raw)
            if not isinstance(previous_quiz, list):
                previous_quiz = []
        except Exception:
            previous_quiz = []
        print("REGENERATE:", regenerate)
        print(
            "PREVIOUS QUIZ QUESTIONS:",
            len(previous_quiz)
        )
        #  Prepare context
        chunks = chunk_text(text)[:15]
        context = "\n".join(chunks)
        #  Build prompt
        prompt = build_quiz_prompt(
            context,
            previous_quiz if regenerate else []
        )
        #  Generate quiz
        raw = call_openai(
            prompt,
            max_tokens=700,
            json_mode=True
        )
        print("RAW AI OUTPUT:", raw)
        #  Parse JSON
        start = raw.find("{")
        end = raw.rfind("}") + 1
        if start == -1 or end <= start:
            return jsonify({
                "quiz": [],
                "error": "Invalid JSON returned by AI"
            }), 500
        clean = raw[start:end]
        data = json.loads(clean)
        quiz = data.get("quiz")
        if not quiz or not isinstance(quiz, list):
            return jsonify({
                "quiz": [],
                "error": "Invalid quiz format from AI"
            }), 500
        return jsonify({
            "quiz": quiz
        })
    except Exception as e:
        print("GENERATE QUIZ ERROR:", str(e))
        return jsonify({
            "quiz": [],
            "error": str(e)
        }), 500
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass

# ========================================================================
# RUN
# ========================================================================
if __name__ == "__main__":
   
    app.run(
    host=os.getenv("FLASK_HOST", "0.0.0.0"),
    port=int(os.getenv("FLASK_PORT", "5000"))
)